"""Shared framework for enterprise DOCX generation (IEEE/ISO style)."""
from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

PROJECT_NAME = "Volunteer Crisis Management Platform"
PROJECT_NAME_FA = "پناه (Panah) — پلتفرم مدیریت بحران و هماهنگی داوطلبان"
VENDOR = "Investica Group"
DOC_VERSION = "1.0.0"
DOC_DATE = date(2026, 7, 16)
CLASSIFICATION = "Internal — Enterprise / Government Deliverable"


def _set_run_font(run, name="Calibri", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def _add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def _add_toc_field(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr)
    run._r.append(fld_char_separate)
    run._r.append(fld_char_end)


class DocBuilder:
    def __init__(
        self,
        title: str,
        doc_id: str,
        subtitle: str = "",
        standard: str = "IEEE / ISO aligned",
    ):
        self.title = title
        self.doc_id = doc_id
        self.subtitle = subtitle
        self.standard = standard
        self.doc = Document()
        self._configure_page()
        self._configure_styles()
        self._section_counter = [0, 0, 0]

    def _configure_page(self):
        section = self.doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.2)
        section.different_first_page_header_footer = True

        header = section.header
        hp = header.paragraphs[0]
        hp.clear()
        run = hp.add_run(f"{PROJECT_NAME}  |  {self.doc_id}  |  v{DOC_VERSION}")
        _set_run_font(run, size=9, color=RGBColor(0x47, 0x55, 0x69))
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.clear()
        run = fp.add_run(f"{CLASSIFICATION}  •  {VENDOR}  •  Page ")
        _set_run_font(run, size=9, color=RGBColor(0x64, 0x74, 0x8B))
        _add_page_number(fp)
        run2 = fp.add_run(f"  •  {DOC_DATE.isoformat()}")
        _set_run_font(run2, size=9, color=RGBColor(0x64, 0x74, 0x8B))
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _configure_styles(self):
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(8)
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

        for level, size in ((1, 16), (2, 13), (3, 12)):
            style = styles[f"Heading {level}"]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            style.paragraph_format.space_before = Pt(14 if level == 1 else 10)
            style.paragraph_format.space_after = Pt(6)

    def add_cover(self):
        for _ in range(3):
            self.doc.add_paragraph()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(PROJECT_NAME_FA)
        _set_run_font(r, size=14, bold=True, color=RGBColor(0x08, 0x91, 0xB2))

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(PROJECT_NAME)
        _set_run_font(r, size=18, bold=True, color=RGBColor(0x0F, 0x17, 0x2A))

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(self.title)
        _set_run_font(r, size=22, bold=True, color=RGBColor(0x1E, 0x3A, 0x5F))

        if self.subtitle:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(self.subtitle)
            _set_run_font(r, size=12, color=RGBColor(0x47, 0x55, 0x69))

        self.doc.add_paragraph()
        meta = [
            ("Document ID", self.doc_id),
            ("Version", DOC_VERSION),
            ("Date", DOC_DATE.isoformat()),
            ("Status", "Approved for Delivery"),
            ("Classification", CLASSIFICATION),
            ("Vendor / Author", VENDOR),
            ("Standards Alignment", self.standard),
        ]
        table = self.doc.add_table(rows=len(meta), cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (k, v) in enumerate(meta):
            table.rows[i].cells[0].text = k
            table.rows[i].cells[1].text = v
            for cell in table.rows[i].cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        _set_run_font(run, size=10)

        self.doc.add_paragraph()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(
            "This document is derived from the production source code, "
            "configuration, and operational architecture of the Panah platform."
        )
        _set_run_font(r, size=9, color=RGBColor(0x64, 0x74, 0x8B))
        self.doc.add_page_break()

    def add_revision_history(self, rows: list[tuple[str, str, str, str]] | None = None):
        self.h1("Document Control")
        self.h2("Revision History")
        if rows is None:
            rows = [
                (DOC_VERSION, DOC_DATE.isoformat(), "Documentation Team", "Initial baseline from source-of-truth codebase"),
            ]
        self.table(["Version", "Date", "Author", "Description"], rows)
        self.h2("Approvals")
        self.table(
            ["Role", "Name", "Signature", "Date"],
            [
                ("Project Sponsor", "Investica Leadership", "", ""),
                ("Technical Architect", "Platform Architecture", "", ""),
                ("Quality Assurance", "QA Lead", "", ""),
            ],
        )
        self.doc.add_page_break()

    def add_toc(self):
        self.h1("Table of Contents")
        p = self.doc.add_paragraph()
        _add_toc_field(p)
        note = self.doc.add_paragraph()
        r = note.add_run(
            "Note: In Microsoft Word, right-click the TOC field and select "
            "“Update Field” to refresh page numbers after opening."
        )
        _set_run_font(r, size=9, color=RGBColor(0x64, 0x74, 0x8B))
        self.doc.add_page_break()

    def h1(self, text: str):
        self._section_counter[0] += 1
        self._section_counter[1] = 0
        self._section_counter[2] = 0
        self.doc.add_heading(f"{self._section_counter[0]}. {text}", level=1)

    def h2(self, text: str):
        self._section_counter[1] += 1
        self._section_counter[2] = 0
        self.doc.add_heading(
            f"{self._section_counter[0]}.{self._section_counter[1]} {text}", level=2
        )

    def h3(self, text: str):
        self._section_counter[2] += 1
        self.doc.add_heading(
            f"{self._section_counter[0]}.{self._section_counter[1]}.{self._section_counter[2]} {text}",
            level=3,
        )

    def p(self, text: str):
        para = self.doc.add_paragraph(text)
        for run in para.runs:
            _set_run_font(run, size=11)
        return para

    def bullets(self, items: Iterable[str]):
        for item in items:
            para = self.doc.add_paragraph(item, style="List Bullet")
            for run in para.runs:
                _set_run_font(run, size=11)

    def numbered(self, items: Iterable[str]):
        for item in items:
            para = self.doc.add_paragraph(item, style="List Number")
            for run in para.runs:
                _set_run_font(run, size=11)

    def table(self, headers: list[str], rows: list[tuple | list]):
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = "Table Grid"
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h
            for para in cell.paragraphs:
                for run in para.runs:
                    _set_run_font(run, size=10, bold=True)
                # shade header
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E2E8F0")
                shading.set(qn("w:val"), "clear")
                cell._tePr = cell._tc.get_or_add_tcPr()
                cell._tc.get_or_add_tcPr().append(shading)
        for r_idx, row in enumerate(rows):
            for c_idx, val in enumerate(row):
                table.rows[r_idx + 1].cells[c_idx].text = str(val)
                for para in table.rows[r_idx + 1].cells[c_idx].paragraphs:
                    for run in para.runs:
                        _set_run_font(run, size=10)
        self.doc.add_paragraph()

    def figure(self, image_path: Path | None, caption: str, width_inches: float = 5.8):
        if image_path and image_path.exists():
            try:
                self.doc.add_picture(str(image_path), width=Inches(width_inches))
                last = self.doc.paragraphs[-1]
                last.alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception:
                self.p(f"[Figure unavailable: {image_path.name}]")
        else:
            self.p(f"[Figure placeholder — diagram to be linked: {caption}]")
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(caption)
        _set_run_font(r, size=9, bold=True, color=RGBColor(0x33, 0x41, 0x55))

    def callout(self, title: str, body: str):
        p = self.doc.add_paragraph()
        r = p.add_run(f"{title}: ")
        _set_run_font(r, size=11, bold=True, color=RGBColor(0x08, 0x91, 0xB2))
        r2 = p.add_run(body)
        _set_run_font(r2, size=11)

    def add_references(self, refs: list[str] | None = None):
        self.h1("References")
        default = [
            "IEEE 29148:2018 — Systems and software engineering — Life cycle processes — Requirements engineering",
            "ISO/IEC/IEEE 42010:2022 — Software, systems and enterprise — Architecture description",
            "UML 2.5 — Unified Modeling Language",
            "BPMN 2.0 — Business Process Model and Notation",
            "C4 Model — Context, Containers, Components, Code (Simon Brown)",
            "OWASP ASVS 4.0 — Application Security Verification Standard",
            "OWASP Top 10 (2021)",
            "REST API Design Rulebook / OpenAPI 3 Specification",
            "Django 5.2 Documentation / Django REST Framework",
            "React 19 Documentation / Material UI v6",
            "PostgreSQL 17 Documentation",
            "Panah Platform Source Code Repository (current workspace)",
            "Panah README.md and docker-compose.yml",
        ]
        self.numbered(refs or default)

    def save(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        return path
