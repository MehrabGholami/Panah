# -*- coding: utf-8 -*-
"""Shared DOCX helpers for RTL Persian IEEE-style SRS."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

FONT = "B Nazanin"
FONT_FALLBACK = "Tahoma"
ACCENT = RGBColor(0x08, 0x91, 0xB2)
DARK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x47, 0x55, 0x69)


def set_run_font(run, size=12, bold=False, color=None, name=FONT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def rtl_paragraph(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def add_toc_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(sep)
    run._r.append(end)


def shade_cell(cell, fill="E0F2FE"):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc.append(shd)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    for style_name, size, bold in (
        ("Heading 1", 18, True),
        ("Heading 2", 15, True),
        ("Heading 3", 13, True),
    ):
        st = styles[style_name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = DARK
        st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    header = section.header.paragraphs[0]
    rtl_paragraph(header)
    hr = header.add_run("پناه (Panah) — Software Requirements Specification | Version 1.0 | IEEE 29148")
    set_run_font(hr, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    rtl_paragraph(footer)
    fr = footer.add_run("طبقه‌بندی: داخلی / قابل ارائه به کارفرما  |  صفحه ")
    set_run_font(fr, size=9, color=MUTED)
    add_page_number(footer)
    fr2 = footer.add_run("  |  Investica Group")
    set_run_font(fr2, size=9, color=MUTED)


def h1(doc, text):
    p = doc.add_heading(text, level=1)
    rtl_paragraph(p)
    for run in p.runs:
        set_run_font(run, size=18, bold=True, color=DARK)
    return p


def h2(doc, text):
    p = doc.add_heading(text, level=2)
    rtl_paragraph(p)
    for run in p.runs:
        set_run_font(run, size=15, bold=True, color=DARK)
    return p


def h3(doc, text):
    p = doc.add_heading(text, level=3)
    rtl_paragraph(p)
    for run in p.runs:
        set_run_font(run, size=13, bold=True, color=DARK)
    return p


def para(doc, text, size=12, bold=False, space_after=8):
    p = doc.add_paragraph()
    rtl_paragraph(p)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=DARK)
    return p


def note(doc, text):
    p = doc.add_paragraph()
    rtl_paragraph(p)
    run = p.add_run("یادداشت: ")
    set_run_font(run, size=11, bold=True, color=ACCENT)
    run2 = p.add_run(text)
    set_run_font(run2, size=11, color=MUTED)
    return p


def tip(doc, text):
    p = doc.add_paragraph()
    rtl_paragraph(p)
    run = p.add_run("نکته: ")
    set_run_font(run, size=11, bold=True, color=RGBColor(0x05, 0x96, 0x69))
    run2 = p.add_run(text)
    set_run_font(run2, size=11, color=MUTED)
    return p


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        rtl_paragraph(p)
        r = p.add_run(h)
        set_run_font(r, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade_cell(cell, "0E7490")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            rtl_paragraph(p)
            r = p.add_run(str(val))
            set_run_font(r, size=10, color=DARK)
            if ri % 2 == 1:
                shade_cell(cell, "F0FDFA")
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def caption(doc, text):
    p = doc.add_paragraph()
    rtl_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, size=10, bold=True, color=MUTED)
    return p


def add_image(doc, path: Path, width_inches=5.8, caption_text: str | None = None):
    if path.exists():
        doc.add_picture(str(path), width=Inches(width_inches))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption_text:
        caption(doc, caption_text)


def page_break(doc):
    doc.add_page_break()
