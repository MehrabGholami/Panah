from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(r"c:\Users\Mehrab\Desktop\مدیریت بحران")
OUT = ROOT / "P3.docx"
ASSETS = ROOT / "P3_assets"
ASSETS.mkdir(exist_ok=True)
FF = next((Path(x) for x in (r"C:\Windows\Fonts\arialbd.ttf", r"C:\Windows\Fonts\arial.ttf") if Path(x).exists()), None)


def f(size):
    return ImageFont.truetype(str(FF), size) if FF else ImageFont.load_default()


def block(d, r, text, fill, size=17):
    d.rounded_rectangle(r, radius=16, fill=fill, outline="#334155", width=2)
    y = r[1] + 18
    for line in text.split("\n"):
        w = d.textbbox((0, 0), line, font=f(size))[2]
        d.text((r[0] + (r[2]-r[0]-w)/2, y), line, font=f(size), fill="#0F172A")
        y += d.textbbox((0, 0), line, font=f(size))[3] + 5


def arrow(d, a, b):
    d.line([a, b], fill="#475569", width=4)
    x, y = b
    d.polygon([(x, y), (x-12, y-7), (x-12, y+7)], fill="#475569")


def make_diagrams():
    # Target architecture
    im = Image.new("RGB", (1600, 900), "#F8FAFC"); d = ImageDraw.Draw(im)
    d.text((45, 30), "Phase 3 Target Architecture — API First Platform", font=f(30), fill="#0F172A")
    items = [
        ((55,330,275,470), "React SPA\nTailwind + shadcn/ui\nZustand + Query", "#DBEAFE"),
        ((390,330,610,470), "Nginx Gateway\nTLS • Routing\nRate Limits", "#FEF3C7"),
        ((725,250,1030,550), "FastAPI\nOpenAPI • JWT • RBAC\nValidation • Error Model", "#CFFAFE"),
        ((1150,80,1500,210), "PostgreSQL 17\nPostGIS • Transactions\nAlembic", "#FCE7F3"),
        ((1150,310,1500,440), "Redis\nCache • OTP • Limits\nTask Broker", "#FEE2E2"),
        ((1150,550,1500,680), "Worker\nNotifications • Retries\nScheduled Tasks", "#EDE9FE"),
        ((1150,750,1500,860), "Email / SMS\nExternal Provider", "#FFEDD5"),
    ]
    for r,t,c in items: block(d,r,t,c)
    arrow(d,(275,400),(390,400)); arrow(d,(610,400),(725,400)); arrow(d,(1030,350),(1150,145))
    arrow(d,(1030,400),(1150,375)); arrow(d,(1325,440),(1325,550)); arrow(d,(1325,680),(1325,750))
    im.save(ASSETS/"target_architecture.png")
    # API lifecycle
    im = Image.new("RGB", (1600, 760), "#FFFFFF"); d = ImageDraw.Draw(im)
    d.text((45, 30), "API Request Lifecycle and Cross-Cutting Controls", font=f(30), fill="#0F172A")
    names = ["Request", "Gateway", "JWT / RBAC", "Pydantic\nValidation", "Service\nTransaction", "Repository", "Response /\nAudit"]
    cols = ["#DBEAFE","#FEF3C7","#EDE9FE","#DCFCE7","#CFFAFE","#FCE7F3","#FFEDD5"]
    x=35
    for i,n in enumerate(names):
        block(d,(x,280,x+180,410),n,cols[i])
        if i<6: arrow(d,(x+180,345),(x+205,345))
        x+=225
    d.text((55,560), "Every response includes a request ID. Exceptions are mapped to a stable RFC 7807-inspired error format.", font=f(18), fill="#334155")
    im.save(ASSETS/"api_lifecycle.png")
    # security
    im = Image.new("RGB", (1600, 850), "#F8FAFC"); d = ImageDraw.Draw(im)
    d.text((45,30),"Authentication, Authorization and Abuse Protection",font=f(30),fill="#0F172A")
    cards=[
        ((80,170,410,350),"OTP / 2FA\nexpiry • attempt cap\nreplay prevention","#DBEAFE"),
        ((510,170,840,350),"JWT HS256\nshort access lifetime\nrotating refresh token","#DCFCE7"),
        ((940,170,1270,350),"RBAC Policy\nadmin • coordinator\nvolunteer","#EDE9FE"),
        ((80,490,410,670),"Rate Limiting\nRedis counters\nroute-aware thresholds","#FEF3C7"),
        ((510,490,840,670),"Input Safety\nPydantic schemas\nallow-listed filters","#FCE7F3"),
        ((940,490,1270,670),"Audit Trail\nactor • action\nresource • request ID","#FFEDD5")]
    for r,t,c in cards:block(d,r,t,c,19)
    im.save(ASSETS/"security_controls.png")
    # geo score
    im = Image.new("RGB",(1600,850),"#FFFFFF");d=ImageDraw.Draw(im)
    d.text((45,30),"Geospatial Search and Explainable Assignment Score",font=f(30),fill="#0F172A")
    block(d,(60,250,330,410),"Mission Area\nPoint / Polygon\nPostGIS Geography","#DBEAFE")
    block(d,(440,250,710,410),"Volunteer Location\nConsent + Timestamp\nGIST Index","#DCFCE7")
    block(d,(820,250,1090,410),"Candidate Filter\nST_DWithin\nAvailability + Status","#FEF3C7")
    block(d,(1200,210,1530,450),"Ranked Result\n45% Skill\n25% Availability\n20% Distance\n10% Priority","#EDE9FE")
    arrow(d,(330,330),(440,330));arrow(d,(710,330),(820,330));arrow(d,(1090,330),(1200,330))
    d.text((60,610),"The coordinator receives factors and remains accountable for the final assignment decision.",font=f(20),fill="#334155")
    im.save(ASSETS/"spatial_assignment.png")
    # delivery pipeline
    im=Image.new("RGB",(1600,800),"#F8FAFC");d=ImageDraw.Draw(im)
    d.text((45,30),"Quality Gates and Repeatable Development Delivery",font=f(30),fill="#0F172A")
    stages=["Format\nLint","Unit\nTests","Integration\nPostgres + Redis","OpenAPI\nContract","Security\nTests","Container\nSmoke Test","Release\nCandidate"]
    x=30
    for i,s in enumerate(stages):
        block(d,(x,300,x+180,435),s,["#DBEAFE","#DCFCE7","#CFFAFE","#FEF3C7","#FCE7F3","#EDE9FE","#FFEDD5"][i])
        if i<len(stages)-1:arrow(d,(x+180,368),(x+205,368))
        x+=225
    im.save(ASSETS/"quality_pipeline.png")


def rtl(p, alignment=WD_ALIGN_PARAGRAPH.RIGHT):
    p.alignment=alignment; pr=p._p.get_or_add_pPr(); e=OxmlElement("w:bidi");e.set(qn("w:val"),"1");pr.append(e);return p
def shade(cell, color):
    pr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),color);pr.append(e)
def page(p):
    r=p.add_run();b=OxmlElement("w:fldChar");b.set(qn("w:fldCharType"),"begin");i=OxmlElement("w:instrText");i.set(qn("xml:space"),"preserve");i.text=" PAGE ";e=OxmlElement("w:fldChar");e.set(qn("w:fldCharType"),"end");r._r.extend([b,i,e])


class Doc:
    def __init__(self):
        self.d=Document();self.n=0
        s=self.d.sections[0];s.top_margin=Cm(2);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2);s.right_margin=Cm(2)
        for style,size,col in [("Normal",12,"000000"),("Title",25,"0F172A"),("Heading 1",18,"0B5E8E"),("Heading 2",15,"1D4ED8")]:
            st=self.d.styles[style];st.font.name="B Nazanin";st._element.rPr.rFonts.set(qn("w:eastAsia"),"B Nazanin");st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(col)
        h=s.header.paragraphs[0];h.text="سامانه مدیریت داوطلبان بحران | سند فاز سوم — API و توسعه اجرایی";rtl(h);h.runs[0].font.size=Pt(9)
        ft=s.footer.paragraphs[0];rtl(ft,WD_ALIGN_PARAGRAPH.CENTER);ft.add_run("P3 | نسخه 1.0 | صفحه ");page(ft)
    def h1(self,t):self.n+=1;p=self.d.add_heading(f"{self.n}. {t}",1);rtl(p)
    def h2(self,t):p=self.d.add_heading(t,2);rtl(p)
    def p(self,t,b=False):
        p=rtl(self.d.add_paragraph());r=p.add_run(t);r.bold=b;p.paragraph_format.space_after=Pt(7)
    def bullets(self,items):
        for x in items:p=rtl(self.d.add_paragraph(style="List Bullet"));p.add_run(x)
    def table(self,heads,rows):
        t=self.d.add_table(rows=1,cols=len(heads));t.style="Light Shading Accent 1";t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for c,v in zip(t.rows[0].cells,heads):
            c.text=v;shade(c,"0B5E8E");c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
            for p in c.paragraphs:
                rtl(p,WD_ALIGN_PARAGRAPH.CENTER)
                for r in p.runs:r.font.color.rgb=RGBColor(255,255,255);r.bold=True
        for row in rows:
            cells=t.add_row().cells
            for c,v in zip(cells,row):
                c.text=str(v);c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
                for p in c.paragraphs:rtl(p)
        self.d.add_paragraph()
    def fig(self,name,caption):
        p=self.d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(ASSETS/name),width=Cm(16.4))
        c=rtl(self.d.add_paragraph("شکل — "+caption),WD_ALIGN_PARAGRAPH.CENTER);c.runs[0].italic=True;c.runs[0].font.size=Pt(10)
    def br(self):self.d.add_page_break()
    def toc(self):
        p=rtl(self.d.add_paragraph());r=p.add_run();b=OxmlElement("w:fldChar");b.set(qn("w:fldCharType"),"begin");i=OxmlElement("w:instrText");i.set(qn("xml:space"),"preserve");i.text=' TOC \\o "1-3" \\h \\z \\u ';e=OxmlElement("w:fldChar");e.set(qn("w:fldCharType"),"end");r._r.extend([b,i,e])


def build():
    make_diagrams();d=Doc()
    for _ in range(5):d.d.add_paragraph()
    p=d.d.add_heading("سند معماری اجرایی، قرارداد API و بسته توسعه",0);rtl(p,WD_ALIGN_PARAGRAPH.CENTER)
    p=rtl(d.d.add_paragraph("فاز سوم سامانه مدیریت داوطلبان بحران (پناه)"),WD_ALIGN_PARAGRAPH.CENTER);p.runs[0].font.size=Pt(18);p.runs[0].font.color.rgb=RGBColor(29,78,216)
    d.d.add_paragraph();d.table(["مقدار","مشخصه"],[["P3","شناسه سند"],["1.0","نسخه"],["1405/04/26","تاریخ"],["API، بک‌اند، امنیت، مکان و داشبورد","دامنه"],["تیم توسعه و کارفرما","مخاطبان"]])
    d.p("این سند، نقشه اجرایی ساخت نسخه اولیه قابل استفاده سامانه و مبنای قرارداد مشترک تیم‌های بک‌اند، فرانت‌اند، آزمون، عملیات و کارفرما است.")
    d.br();p=d.d.add_heading("فهرست مطالب",1);rtl(p);d.toc();d.br()
    d.h1("کنترل سند، تصمیم معماری و دامنه")
    d.table(["شرح","نسخه","تاریخ"],[["ایجاد نخستین نسخه معماری اجرایی و قرارداد API","1.0","1405/04/26"]])
    d.p("درخواست فاز سوم، FastAPI، SQLAlchemy، Alembic، PostGIS، React، Tailwind و shadcn/ui را به‌عنوان پشته هدف معرفی می‌کند. پیاده‌سازی موجود در مخزن از Django REST Framework و React/MUI استفاده می‌کند. ازاین‌رو، این سند معماری هدف را تعریف می‌کند و نباید آن را تأیید تبدیل انجام‌شده کد فعلی تلقی کرد.")
    d.bullets(["خروجی هدف: API عملیاتی، داشبورد نقش‌محور، امنیت OTP/JWT/RBAC، منطق تخصیص و بسته Docker.","مرزها: داده‌های عملیاتی و قراردادها باید بدون شکستن کاربران فعلی قابل انتقال باشند.","اصل حاکم: هیچ دو سازوکار migration نباید هم‌زمان مالک تغییر یک جدول تولیدی باشند."])
    d.h1("معماری هدف و مرزبندی لایه‌ها")
    d.p("معماری API-first از لایه ارائه مستقل است. Router فقط پروتکل HTTP، احراز هویت و تبدیل DTO را انجام می‌دهد؛ Service قواعد کاربردی و مرز تراکنش را مالک است؛ Repository تنها به SQLAlchemy و ذخیره‌سازی وابسته است. این تفکیک تست‌پذیری و امکان تعویض راهبرد ذخیره را افزایش می‌دهد.")
    d.fig("target_architecture.png","معماری هدف کانتینری؛ تمام برچسب‌های داخل تصویر انگلیسی هستند.")
    d.h2("مسئولیت سرویس‌ها")
    d.table(["نمونه مسئولیت","لایه"],[["Route، dependency injection، schema validation، status code","API Router"],["حالت‌های مجاز مأموریت، OTP، RBAC و امتیازدهی تخصیص","Domain / Service"],["Query، قفل سطری، flush و persistence","Repository"],["ORM mapping و Alembic revisions","Data Access"],["ارسال اعلان با retry و idempotency","Worker"]])
    d.h1("قرارداد داده و استاندارد API")
    d.p("تمام endpointها زیر /api/v1 قرار می‌گیرند. تغییر افزایشی و سازگار در v1 ممکن است؛ هر تغییر ناسازگار، حذف فیلد، تغییر معنای وضعیت یا رفتار امنیتی نیازمند v2 و بازه deprecation است. OpenAPI همان مرجع قابل تولید برای Swagger UI، SDK و contract test خواهد بود.")
    d.fig("api_lifecycle.png","چرخه استاندارد هر درخواست API و کنترل‌های مشترک.")
    d.h2("قواعد Request و Response")
    d.bullets(["Content-Type درخواست‌ها application/json است؛ فایل‌ها فقط از multipart/form-data و با سقف اندازه مشخص پذیرفته می‌شوند.","شناسه‌ها UUID هستند و زمان‌ها ISO-8601 با timezone؛ تاریخ‌های محلی فقط در لایه ارائه تبدیل می‌شوند.","لیست‌ها data، meta و links دارند. page پیش‌فرض 1، page_size پیش‌فرض 20 و حداکثر 100 است.","فیلترها به شکل filter[status] و بازه‌ها به شکل filter[start_at_gte] هستند. sort با -created_at برای نزولی نوشته می‌شود.","Request-ID ورودی پذیرفته یا تولید می‌شود و در پاسخ/لاگ/خطا بازگردانده می‌شود."])
    d.table(["کاربرد","کد HTTP"],[["ایجاد موفق منبع","201 Created"],["درخواست معتبر و پردازش‌شده","200 OK"],["نبود یا انقضای توکن","401 Unauthorized"],["نقش یا مالکیت ناکافی","403 Forbidden"],["منبع ناموجود یا پنهان‌شده","404 Not Found"],["تعارض درخواست تکراری یا ظرفیت","409 Conflict"],["اعتبارسنجی ساختار/فیلد","422 Unprocessable Entity"],["محدودیت نرخ","429 Too Many Requests"],["خطای پیش‌بینی‌نشده","500 با شناسه پیگیری، بدون افشای جزئیات"]])
    d.h2("Endpointهای پایه")
    d.table(["نمونه Endpoint","عملیات"],[["POST /v1/auth/otp/request","درخواست چالش OTP"],["POST /v1/auth/otp/verify","اعتبارسنجی OTP و صدور token"],["GET, POST /v1/volunteers","فهرست/ایجاد داوطلب"],["GET /v1/volunteers/nearby","جست‌وجوی مکانی محدودشده"],["GET, POST /v1/missions","فهرست/ایجاد مأموریت"],["POST /v1/missions/{id}/apply","ثبت درخواست یکتا"],["POST /v1/applications/{id}/approve","تأیید تراکنشی و ایجاد تخصیص"],["GET, POST /v1/missions/{id}/shifts","مدیریت شیفت"],["GET, POST /v1/missions/{id}/reports","گزارش مأموریت"]])
    d.h1("بک‌اند FastAPI، SQLAlchemy و Alembic")
    d.p("مدل‌های SQLAlchemy باید از مدل‌های Pydantic جدا بمانند. DTO خروجی هرگز فیلدهای داخلی مانند password_hash، OTP hash، refresh token یا داده پزشکی خارج از سطح مجاز را نشت نمی‌دهد. Unit of Work تراکنش تغییرهای چندجدولی مانند تأیید درخواست را کنترل می‌کند.")
    d.h2("Repository و Service Layer")
    d.table(["قاعده","پیامد اجرایی"],[["Router فاقد منطق کسب‌وکار","تست واحد سرویس بدون HTTP ممکن می‌شود"],["Repository فاقد تصمیم مجوز","کنترل نقش در policy/service یکپارچه می‌ماند"],["تراکنش در service","عملیات approve شامل application، capacity، assignment و outbox اتمیک است"],["Idempotency key در عملیات حساس","retry شبکه درخواست تکراری ایجاد نمی‌کند"],["Alembic revision کوچک و قابل بازبینی","استقرار و بازگشت کنترل‌شده می‌شود"]])
    d.p("برای تغییرهای خطرناک schema، الگوی expand/contract پیشنهاد می‌شود: ابتدا ستون/جدول جدید سازگار افزوده شود؛ کد دوگانه‌خوان/دوگانه‌نویس موقت مستقر شود؛ داده backfill شود؛ پس از اندازه‌گیری و پایان دوره سازگاری، ساختار قدیمی حذف شود.")
    d.h1("امنیت: OTP، JWT، RBAC و کنترل‌های پایه")
    d.fig("security_controls.png","لایه‌های کنترل هویت، مجوز، ورودی، نرخ و ممیزی.")
    d.h2("جریان هویت")
    d.p("OTP باید فقط به‌صورت hash ذخیره شود، زمان انقضای کوتاه داشته باشد، بعد از مصرف invalid شود و شمارش تلاش/ارسال آن با Redis محدود گردد. پس از تأیید، access token کوتاه‌عمر و refresh token چرخشی صادر می‌شود. HS256 تنها با secret تصادفی بلند در secret manager، iss/aud/exp/jti و برنامه چرخش کلید قابل استفاده است.")
    d.h2("سیاست RBAC و مالکیت")
    d.table(["نمونه مجوز","نقش/قاعده"],[["ایجاد یا بستن رخداد","Admin"],["ویرایش مأموریت و بررسی درخواست","Coordinator مالک مأموریت یا Admin"],["ثبت درخواست برای مأموریت منتشرشده","Volunteer فعال"],["مشاهده اطلاعات حساس داوطلب","Coordinator مرتبط با حداقل فیلد لازم"],["مشاهده audit کامل","Admin یا نقش ممیزی مصوب"]])
    d.bullets(["bcrypt با cost قابل تنظیم و rehash تدریجی استفاده می‌شود.","Rate limit برای OTP، login، refresh، upload و search از یک آستانه عمومی سخت‌گیرانه‌تر است.","CORS allow-list، security headers، validation طول/نوع/محدوده، و جلوگیری از IDOR اجباری است.","خطاها به کاربر پیام امن می‌دهند اما جزئیات فنی با request_id در لاگ نگهداری می‌شود."])
    d.h1("قابلیت مکانی و منطق تخصیص")
    d.fig("spatial_assignment.png","جست‌وجوی PostGIS و امتیازدهی شفاف پیشنهاد تخصیص.")
    d.p("PostGIS برای search دقیق بر مبنای geography(Point, 4326) یا محدوده عملیاتی استفاده می‌شود. ستون location_geog با GIST index، موقعیت داوطلب با timestamp و وضعیت رضایت، و queryهای ST_DWithin/KNN فقط پس از کنترل مجوز و هدف پردازش قابل استفاده‌اند.")
    d.h2("الگوریتم اولیه تخصیص")
    d.table(["وزن اولیه","عامل"],[["45%","تطابق مهارت‌های الزامی و ترجیحی"],["25%","آمادگی و هم‌پوشانی زمانی با شیفت"],["20%","فاصله نرمال‌شده از محدوده مأموریت"],["10%","اولویت عملیاتی/تجربه"],["قید سخت","فعال بودن داوطلب، ظرفیت، رضایت مکانی و نبود تخصیص متعارض"]])
    d.p("خروجی الگوریتم، پیشنهاد رتبه‌بندی‌شده و عوامل امتیاز است؛ تصمیم نهایی برای عملیات بحرانی باید در اختیار هماهنگ‌کننده بماند. هر انتخاب یا override باید در audit log ثبت شود.")
    d.h1("فرانت‌اند: داشبوردها و جریان‌های بحرانی")
    d.p("رابط کاربری بحران باید سریع، قابل فهم در فشار زمانی و مقاوم در برابر شبکه ضعیف باشد. Tailwind و shadcn/ui اجزای قابل ترکیب را فراهم می‌کنند؛ TanStack Query فقط server state و cache را مدیریت می‌کند؛ Zustand برای وضعیت کوتاه‌عمر UI، فیلترهای محلی و draftهای محدود است.")
    d.table(["جریان اصلی","پنل"],[["مدیریت کاربران، نقش، رخداد، سیاست و گزارش سراسری","Admin"],["صندوق درخواست، ظرفیت، تخصیص، شیفت و گزارش مأموریت‌های خود","Coordinator"],["تکمیل پروفایل/مهارت، مشاهده مأموریت، درخواست، پاسخ به تخصیص و گزارش","Volunteer"]])
    d.bullets(["نمایش وضعیت‌ها با متن، آیکون و رنگ؛ رنگ هرگز تنها حامل معنی نیست.","تأییدهای مهم شامل نام مأموریت، اثر تغییر و امکان لغو امن‌اند.","خطای شبکه قابل تشخیص، retry کنترل‌شده و draft محلی برای فرم‌های طولانی دارد.","دسترسی موبایل‌محور، RTL، خوانایی بالا و حداقل کلیک برای عملیات اضطراری الزام است."])
    d.h1("تست‌پذیری، مشاهده‌پذیری و کیفیت فنی")
    d.fig("quality_pipeline.png","دروازه‌های کیفیت از تغییر کد تا بسته استقرار.")
    d.h2("راهبرد آزمون")
    d.table(["نمونه","سطح آزمون"],[["محاسبه score، قوانین انتقال وضعیت و policy نقش","Unit"],["تعارض unique application، rollback تراکنش و query مکانی","Integration با PostgreSQL/PostGIS و Redis"],["نمونه OpenAPI، envelope خطا و backward compatibility","Contract"],["ثبت‌نام تا تأیید مأموریت و اعلان","End-to-End"],["انقضای token، rate limit، IDOR و input fuzzing","Security"]])
    d.p("لاگ‌ها باید structured JSON و دارای request_id، actor_id، route، latency و نتیجه باشند. PII و token در log ممنوع است. معیارها شامل نرخ خطا، زمان پاسخ، backlog صف، نرخ OTP، query کند و درصد موفقیت اعلان هستند.")
    d.h1("بسته اجرایی Docker و آماده‌سازی محیط")
    d.p("Docker Compose محیط یکسان برای API، worker، PostgreSQL 17 با PostGIS، Redis، frontend و Nginx فراهم می‌کند. .env.example فقط نام متغیر و مقدار غیرحساس دارد. migrationها در pipeline کنترل‌شده اجرا می‌شوند؛ اجرای خودکار schema migration در production بدون gate عملیاتی توصیه نمی‌شود.")
    d.table(["معیار پذیرش","قلم تحویلی"],[["Swagger/OpenAPI دارای مثال و کد خطای پایدار","مستند API"],["اجرای یک فرمانی محیط توسعه و seed idempotent","Docker Compose + راهنما"],["تغییر schema نسخه‌بندی‌شده و آزمون‌شده","Alembic migrations"],["سناریوی درخواست تا تخصیص قابل تست","API و تست‌های integration"],["پنل‌های نقش‌محور با مسیرهای اصلی","React dashboards"],["کنترل OTP/JWT/RBAC و گزارش ممیزی","امنیت پایه"]])
    d.h1("برنامه اجرا، ریسک‌ها و معیار پذیرش")
    d.table(["خروجی","مرحله"],[["تثبیت OpenAPI، مدل داده هدف و مرزبندی migration","۱. طراحی قرارداد"],["هویت، RBAC، skeleton سرویس و migrations","۲. زیرساخت بک‌اند"],["داوطلب/مهارت/مأموریت/درخواست/تخصیص","۳. عملیات اصلی"],["PostGIS، شیفت، گزارش، اعلان و audit","۴. قابلیت‌های تکمیلی"],["داشبوردها، آزمون E2E، hardening و بسته توسعه","۵. پذیرش"]])
    d.bullets(["ریسک اصلی: دو مالک schema (Django migration و Alembic). راهکار: مالکیت واحد جدول یا مرز سرویس جدا.","ریسک اصلی: اتکا به تخصیص خودکار. راهکار: پیشنهاد explainable با تأیید انسانی.","ریسک اصلی: نشت مکان و PII. راهکار: رضایت، حداقل‌سازی، retention و RBAC.","ریسک اصلی: تناقض API. راهکار: OpenAPI contract testing و version policy."])
    d.h1("مراجع و اسناد همراه")
    d.bullets(["P31.md — معماری فنی، قراردادها، دیاگرام‌های Mermaid و برنامه مهاجرت.","OpenAPI Specification 3.1، FastAPI، SQLAlchemy 2.0، Alembic، PostGIS Documentation.","OWASP ASVS و OWASP API Security Top 10.","RFC 7807 Problem Details، JWT RFC 7519، ISO/IEC 27001 و ISO/IEC 25010."])
    d.p("پایان سند.",True);d.d.save(OUT);print("P3.docx created")

if __name__=="__main__":
    build()
